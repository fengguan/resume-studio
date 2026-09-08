import copy
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from pypdf import PdfReader

from .parsing import document_paragraphs


class RenderError(RuntimeError):
    pass


def paragraph_map(doc):
    return {
        f"b{i:03d}": p
        for i, (_, p) in enumerate(
            item for item in document_paragraphs(doc) if item[1].text.strip()
        )
    }


def replace_text(para, text):
    """Deliberately rebuild only an approved text block, keeping paragraph and base run style."""
    base = max(para.runs, key=lambda r: len(r.text), default=None)
    props = copy.deepcopy(base._r.rPr) if base is not None and base._r.rPr is not None else None
    para.clear()
    run = para.add_run(text)
    if props is not None:
        run._r.insert(0, props)


def anchor_runs(para, quote, highlight=False):
    """Split simple text runs at exact boundaries without losing original run formatting."""
    start = para.text.find(quote) if quote else -1
    if start < 0 or para.hyperlinks or "".join(r.text for r in para.runs) != para.text:
        runs = [r for r in para.runs if r.text]
        if highlight:
            for r in runs:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return runs
    end = start + len(quote)
    offset = 0
    selected = []
    for run in list(para.runs):
        original = run.text
        run_end = offset + len(original)
        if offset < end and run_end > start:
            a, b = max(0, start - offset), min(len(original), end - offset)
            props = copy.deepcopy(run._r.rPr)
            tail = run._r
            segments = [(original[:a], False), (original[a:b], True), (original[b:], False)]
            for text, included in segments:
                if not text:
                    continue
                element = copy.deepcopy(run._r)
                new = Run(element, para)
                new.clear()
                if props is not None and new._r.rPr is None:
                    new._r.insert(0, copy.deepcopy(props))
                new.text = text
                tail.addnext(element)
                tail = element
                if included:
                    selected.append(new)
                    if highlight:
                        new.font.highlight_color = WD_COLOR_INDEX.YELLOW
            para._p.remove(run._r)
        offset = run_end
    return selected


def render_resume(source: bytes, final_text: dict, risks, annotate: bool) -> bytes:
    doc = Document(BytesIO(source))
    paragraphs = paragraph_map(doc)
    locations = {
        f"b{i:03d}": location
        for i, (location, _) in enumerate(
            item for item in document_paragraphs(doc) if item[1].text.strip()
        )
    }
    for block_id, text in final_text.items():
        if block_id not in paragraphs:
            raise RenderError("导出时找不到原文位置。")
        if paragraphs[block_id].text != text:
            replace_text(paragraphs[block_id], text)
    if annotate:
        for risk in risks:
            para = paragraphs[risk.block_id]
            label = {
                "pending": "待核实",
                "reverted": "已拦截并回退",
                "resolved": "已处理并复核",
                "deleted": "已删除",
            }[risk.status]
            anchor = risk.quote if risk.quote in para.text else para.text
            if not para.text.strip():
                # The review-only placeholder is explicitly not a candidate assertion.
                replace_text(para, "[审阅标记：此处内容已删除]")
            runs = anchor_runs(para, anchor, highlight=risk.status == "pending")
            if not locations[risk.block_id].startswith("body/"):
                # Word does not support comment anchors in headers/footers. Highlight
                # the original text and attach its comment to a review-only body note.
                note = doc.add_paragraph(
                    f"[审阅提示 {risk.id}：页眉/页脚内容需检查，位置 {locations[risk.block_id]}]"
                )
                runs = note.runs
            if not runs:
                raise RenderError("无法为风险添加 Word 批注，未发布无标注文件。")
            doc.add_comment(
                runs,
                text=(
                    f"{risk.id}｜{label}｜"
                    f"{'原文已有疑点' if risk.origin == 'source' else '改写检查'}\n"
                    f"问题表述：{risk.quote}\n原文：{risk.source_text}\n"
                    f"原因：{risk.reason}\n建议：{risk.suggestion}"
                ),
                author="Resume Studio",
                initials="RS",
            )
    out = BytesIO()
    doc.save(out)
    result = out.getvalue()
    # Reopen and verify all original positions, including now-empty paragraphs.
    reread = {location: p.text for location, p in document_paragraphs(Document(BytesIO(result)))}
    original_locations = {
        f"b{i:03d}": location
        for i, (location, _) in enumerate(
            item for item in document_paragraphs(Document(BytesIO(source))) if item[1].text.strip()
        )
    }
    for block_id, expected in final_text.items():
        actual = reread.get(original_locations[block_id])
        if annotate and not expected.strip() and actual == "[审阅标记：此处内容已删除]":
            continue
        if actual != expected:
            raise RenderError("导出文档重读与最终正文不一致，已停止发布。")
    return result


def check_layout(docx_bytes: bytes, target_pages=2, executable=None):
    program = executable or shutil.which("libreoffice") or shutil.which("soffice")
    if not program:
        return {
            "status": "unverified",
            "pages": None,
            "target_pages": target_pages,
            "message": "未安装 LibreOffice，未验证分页；请用 Word 打开审阅。",
        }
    try:
        with tempfile.TemporaryDirectory(prefix="resume-layout-") as tmp:
            root = Path(tmp)
            source = root / "resume.docx"
            source.write_bytes(docx_bytes)
            subprocess.run(
                [
                    program,
                    f"-env:UserInstallation={(root / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    tmp,
                    str(source),
                ],
                capture_output=True,
                timeout=60,
                check=True,
            )
            pdf = root / "resume.pdf"
            if not pdf.exists():
                raise OSError("Missing rendered PDF")
            reader = PdfReader(pdf)
            pages = len(reader.pages)
            has_text = all((p.extract_text() or "").strip() for p in reader.pages)
            ok = pages <= target_pages and has_text
            return {
                "status": "checked" if ok else "needs_review",
                "pages": pages,
                "target_pages": target_pages,
                "message": f"LibreOffice 渲染 {pages} 页，目标最多 {target_pages} 页。"
                "字体替换、视觉布局和 Word 分页仍需人工检查。",
            }
    except (OSError, subprocess.SubprocessError, ValueError):
        return {
            "status": "unverified",
            "pages": None,
            "target_pages": target_pages,
            "message": "PDF 渲染失败，未验证分页；DOCX 仍可下载审阅。",
        }
