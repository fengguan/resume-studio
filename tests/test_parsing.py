import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from resume_studio.parsing import InputError, parse_job, parse_resume


def test_document_order_dates_and_editable_headline(resume_bytes):
    resume = parse_resume(resume_bytes)
    texts = [b.text for b in resume.blocks]
    assert (
        texts.index("EXAMPLE COMPANY")
        < texts.index("2020 - 2026")
        < texts.index("Project Lead (2020 - 2026)")
    )
    date = next(b for b in resume.blocks if b.text == "2020 - 2026")
    assert not date.private and not date.editable
    assert resume.blocks[2].editable
    assert resume.blocks[1].private and not resume.blocks[1].editable


def test_merged_table_and_nested_content(resume_bytes):
    doc = Document(BytesIO(resume_bytes))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged once"
    table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested content"
    doc.sections[0].header.paragraphs[0].text = "Header content"
    out = BytesIO()
    doc.save(out)
    texts = [b.text for b in parse_resume(out.getvalue()).blocks]
    assert texts.count("Merged once") == 1
    assert "Nested content" in texts and "Header content" in texts


@pytest.mark.parametrize("tag", ["w:ins", "w:txbxContent", "w:sdt"])
def test_unsupported_content_fails_explicitly(resume_bytes, tag):
    doc = Document(BytesIO(resume_bytes))
    doc.paragraphs[3]._p.append(OxmlElement(tag))
    out = BytesIO()
    doc.save(out)
    with pytest.raises(InputError):
        parse_resume(out.getvalue())


@pytest.mark.parametrize("bad", [b"", b"not a docx", b"PKbroken"])
def test_bad_docx(bad):
    with pytest.raises(InputError):
        parse_resume(bad)


@pytest.mark.parametrize("bad", [b"{", b"[]", b"{}", b'{"job": {"title":"A", "company":"B"}}'])
def test_bad_job(bad):
    with pytest.raises(InputError):
        parse_job(bad)


def test_sample_inputs():
    root = Path(__file__).resolve().parents[1] / ".data"
    if not root.exists():
        pytest.skip("Private samples not installed")
    resume = parse_resume(next(root.glob("*.docx")).read_bytes())
    assert any(b.text == "BAYER Crop Science" for b in resume.blocks)
    assert any("PEKING UNIVERSITY" in b.text for b in resume.blocks)
    jobs = [parse_job(p.read_bytes()) for p in root.rglob("*.json")]
    assert len(jobs) == 5
    rocket = next(j for j in jobs if "Rocket" in j.company)
    assert rocket.warnings
    assert "Coinbase" not in json.dumps(rocket.model_dump())
    amgen = next(j for j in jobs if j.company == "Amgen")
    assert len([r for r in amgen.requirements if r.category == "required"]) == 3
