import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from resume_studio.parsing import parse_resume
from resume_studio.pipeline import run_pipeline
from resume_studio.schemas import Audit, Change, Verdict


def make_change(resume_bytes, replacement):
    b = next(b for b in parse_resume(resume_bytes).blocks if "approval rates" in b.text)
    return Change(
        block_id=b.id,
        before=b.text,
        after=b.text.replace("approval rates", replacement),
        evidence_ids=[b.id],
        requirement_ids=["j000"],
        reason="突出影响",
    )


def test_blocked_drift_is_reverted_and_annotated(resume_bytes, job_bytes, tmp_path, fake_model):
    change = make_change(resume_bytes, "approval velocity")
    model = fake_model(changes=[change])
    result = run_pipeline(resume_bytes, job_bytes, model, tmp_path)
    assert result.clean_available and result.audit_complete
    assert result.final_text[change.block_id] == change.before
    assert any(r.status == "reverted" for r in result.risks)
    root = Path(result.output_dir)
    with ZipFile(root / "resume_comparison.docx") as archive:
        xml = archive.read("word/document.xml").decode()
        assert "approval velocity" not in xml
        assert "w:ins " not in xml and "w:del " not in xml
    with ZipFile(root / "resume_review.docx") as archive:
        assert "approval velocity" in archive.read("word/comments.xml").decode()
        assert "approval velocity" not in archive.read("word/document.xml").decode()
    with ZipFile(root / "tailored_resume.docx") as archive:
        assert "word/comments.xml" not in archive.namelist()
    comments = (root / "comments.md").read_text()
    assert "已拦截/回退" in comments
    assert "本次未保留对原文的修改" in comments
    assert not any("person@example.com" in json.dumps(payload) for _, payload in model.payloads)


def test_suspected_claim_highlight_and_resolution(resume_bytes, job_bytes, tmp_path, fake_model):
    change = make_change(resume_bytes, "approval rates")
    change.after = change.before.replace("supporting $10B", "owning $10B")

    def audit(payload):
        texts = {b["id"]: b["text"] for b in payload["candidate"]}
        return Audit(
            verdicts=[
                Verdict(
                    block_id=b,
                    severity="suspected" if "owning" in texts[b] else "clear",
                    origin="rewrite",
                    quote="owning $10B" if "owning" in texts[b] else "",
                    reason="归因可能被加强" if "owning" in texts[b] else "",
                    suggestion="恢复 supporting",
                )
                for b in payload["review_block_ids"]
            ]
        )

    model = fake_model(changes=[change], audit_fn=audit)
    result = run_pipeline(resume_bytes, job_bytes, model, tmp_path)
    assert not result.clean_available
    root = Path(result.output_dir)
    assert not (root / "tailored_resume.docx").exists()
    review = Document(root / "resume_review.docx")
    runs = [r for p in review.paragraphs for r in p.runs if r.font.highlight_color]
    assert "".join(r.text for r in runs) == "owning $10B"
    assert "R-001" in list(review.comments)[0].text
    resolved = run_pipeline(
        resume_bytes,
        job_bytes,
        model,
        tmp_path,
        previous=result,
        edits={change.block_id: change.before},
    )
    assert resolved.clean_available
    assert resolved.run_id != result.run_id
    assert any(r.status == "resolved" for r in resolved.risks)
    assert root.exists() and not (root / "tailored_resume.docx").exists()


def test_audit_failure_does_not_create_clean_resume(resume_bytes, job_bytes, tmp_path, fake_model):
    result = run_pipeline(resume_bytes, job_bytes, fake_model(fail_audit=True), tmp_path)
    assert not result.audit_complete and not result.clean_available
    assert any(r.origin == "system" and r.status == "pending" for r in result.risks)
    assert (Path(result.output_dir) / "resume_review.docx").exists()


def test_missing_audit_coverage_fails_closed(resume_bytes, job_bytes, tmp_path, fake_model):
    result = run_pipeline(
        resume_bytes, job_bytes, fake_model(audit_fn=lambda _: Audit(verdicts=[])), tmp_path
    )
    assert not result.audit_complete and not result.clean_available


def test_semantic_violation_reaudits_after_revert(resume_bytes, job_bytes, tmp_path, fake_model):
    change = make_change(resume_bytes, "approval rates")
    change.after = change.before.replace("supporting $10B", "personally generating $10B")

    def audit(payload):
        texts = {b["id"]: b["text"] for b in payload["candidate"]}
        return Audit(
            verdicts=[
                Verdict(
                    block_id=b,
                    severity="violation" if "personally generating" in texts[b] else "clear",
                    origin="rewrite",
                    quote=texts[b],
                    reason="个人归因无依据",
                    suggestion="回退",
                )
                for b in payload["review_block_ids"]
            ]
        )

    model = fake_model(changes=[change], audit_fn=audit)
    result = run_pipeline(resume_bytes, job_bytes, model, tmp_path)
    assert result.clean_available and result.final_text[change.block_id] == change.before
    assert [s for s, _ in model.payloads].count("audit") == 2


def test_source_question_cannot_be_dismissed_by_reverting(
    resume_bytes, job_bytes, tmp_path, fake_model
):
    doc = Document(BytesIO(resume_bytes))
    doc.add_paragraph("Generative AI Certification — in progress; expected Aug 2020")
    out = BytesIO()
    doc.save(out)
    rb = out.getvalue()
    result = run_pipeline(rb, job_bytes, fake_model(), tmp_path)
    r = next(r for r in result.risks if r.origin == "source")
    revised = run_pipeline(
        rb, job_bytes, fake_model(), tmp_path, previous=result, edits={r.block_id: r.source_text}
    )
    assert not revised.clean_available
    removed = run_pipeline(
        rb, job_bytes, fake_model(), tmp_path, previous=result, edits={r.block_id: ""}
    )
    assert removed.clean_available
    assert any(r.status == "deleted" for r in removed.risks)
    with ZipFile(Path(removed.output_dir) / "tailored_resume.docx") as archive:
        assert "审阅标记" not in archive.read("word/document.xml").decode()


def test_new_user_fact_can_resolve_certificate(resume_bytes, job_bytes, tmp_path, fake_model):
    doc = Document(BytesIO(resume_bytes))
    doc.add_paragraph("Generative AI Certification — in progress; expected Aug 2020")
    out = BytesIO()
    doc.save(out)
    rb = out.getvalue()
    result = run_pipeline(rb, job_bytes, fake_model(), tmp_path)
    r = next(r for r in result.risks if r.origin == "source")
    client = fake_model()
    revised = run_pipeline(
        rb,
        job_bytes,
        client,
        tmp_path,
        previous=result,
        edits={r.block_id: "Generative AI Certification — earned Sep 2020"},
        supplements={r.block_id: "本人已于 2020 年 9 月取得该证书。"},
    )
    assert revised.clean_available
    assert r.block_id in client.payloads[-1][1]["user_facts"]


def test_all_sample_jobs_file_pipeline(tmp_path, fake_model):
    root = Path(__file__).resolve().parents[1] / ".data"
    if not root.exists():
        pytest.skip("Private samples not installed")
    rb = next(root.glob("*.docx")).read_bytes()
    for p in root.rglob("*.json"):
        result = run_pipeline(rb, p.read_bytes(), fake_model(), tmp_path)
        assert result.audit_complete
        assert result.risks  # expired source certification remains visible
        assert not result.clean_available
        assert (Path(result.output_dir) / "resume_review.docx").exists()
