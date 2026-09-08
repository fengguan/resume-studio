import json
import os
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import streamlit as st
from docx import Document

from resume_studio.pipeline import run_pipeline
from resume_studio.review_ui import interactive_review
from resume_studio.schemas import Analysis, Audit, Draft, Mapping, Refinement, Verdict
from resume_studio.workspace import load_result

ROOT = Path(os.environ["RESUME_BROWSER_TEST_ROOT"])


class Mock:
    def __init__(self):
        self.config = SimpleNamespace(provider="openai", model="BROWSER-TEST-DOUBLE")
        self.usage = []

    def generate(self, stage, system, payload, output_type):
        self.usage.append({"stage": stage})
        if stage == "analysis":
            return Analysis(
                positioning="测试",
                priorities=[],
                questions=[],
                mappings=[
                    Mapping(
                        requirement_id=r["id"], status="missing", evidence_ids=[], reason="测试"
                    )
                    for r in payload["job"]["requirements"]
                ],
            )
        if stage == "tailoring":
            return Draft(changes=[])
        if stage == "refine":
            text = payload["current_text"].replace("with experience translating", "translating")
            if "SQL" in payload["feedback"]:
                text += " Expert SQL."
            return Refinement(
                block_id=payload["selected_block_id"],
                action="revise",
                text=text,
                message="已精简这段，并保留原文依据。",
                evidence_ids=[payload["selected_block_id"]],
            )
        return Audit(
            verdicts=[
                Verdict(
                    block_id=k,
                    severity="clear",
                    origin="source",
                    quote="",
                    reason="",
                    suggestion="",
                )
                for k in payload["review_block_ids"]
            ]
        )


st.set_page_config(layout="wide")
st.title("交互审阅 · 自动化测试样本")
if "result" not in st.session_state:
    saved = sorted((ROOT / "outputs").glob("*/result.json"))
    if saved:
        st.session_state["result"] = load_result(saved[-1].parent)
    else:
        doc = Document()
        doc.add_paragraph("TEST CANDIDATE")
        doc.add_paragraph("STRATEGY AND ANALYTICS", style="Heading 1")
        doc.add_paragraph(
            "Strategy leader with experience translating market research into executive recommendations and leading cross-functional projects across scientific and business teams."
        )
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "EXAMPLE COMPANY"
        table.cell(0, 1).text = "2020 - 2026"
        doc.add_paragraph(
            "Led cross-functional teams to redesign the drug-evaluation system, increasing executive approval rates by 50% and supporting $10B in revenue potential."
        )
        doc.add_paragraph("Certification in progress; expected Aug 2026")
        stream = BytesIO()
        doc.save(stream)
        jd = json.dumps(
            {"job": {"title": "Strategy Manager", "company": "Example"}, "skills": ["SQL"]}
        ).encode()
        st.session_state["result"] = run_pipeline(stream.getvalue(), jd, Mock(), ROOT / "outputs")
result = st.session_state["result"]
dirty = interactive_review(result, Mock)
st.caption("DIRTY" if dirty else "SYNCED")
if not dirty:
    st.download_button(
        "下载测试 Word",
        (Path(result.output_dir) / "resume_review.docx").read_bytes(),
        "resume_review.docx",
    )
