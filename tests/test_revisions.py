import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from resume_studio.parsing import document_paragraphs, parse_resume
from resume_studio.rendering import RenderError
from resume_studio.revisions import export_saved_comparison, render_comparison
from resume_studio.schemas import Risk


def accepted_or_rejected(data, accept):
    """Independently apply the standard ins/del operations, then let python-docx read it."""
    doc = Document(BytesIO(data))
    for _, para in document_paragraphs(doc):
        for tag in ("w:ins", "w:del"):
            for node in list(para._p.iter(qn(tag))):
                parent = node.getparent()
                if (tag == "w:ins") == accept:
                    for child in list(node):
                        for text in child.iter(qn("w:delText")):
                            text.tag = qn("w:t")
                        node.addprevious(child)
                parent.remove(node)
    return doc


@pytest.mark.parametrize(
    "replacement",
    [
        "Evaluated 200+ projects, prioritizing 30+ concepts.",
        "",
        "新增文字 & 保留 <符号> 以及 10+ 项目。",
        "Original\twith two  spaces\nand newline",
    ],
)
def test_both_revision_views_match_exact_text(resume_bytes, replacement):
    resume = parse_resume(resume_bytes)
    block = next(b for b in resume.blocks if "Evaluated 200+" in b.text)
    final = {b.id: b.text for b in resume.blocks}
    final[block.id] = replacement
    result = render_comparison(resume_bytes, final)
    paths = {b.id: b.location for b in resume.blocks}
    for accept in (True, False):
        doc = accepted_or_rejected(result, accept)
        values = {path: p.text for path, p in document_paragraphs(doc)}
        for b in resume.blocks:
            assert values[paths[b.id]] == (final[b.id] if accept else b.text)


def test_word_level_markup_preserves_context_and_run_styles(resume_bytes):
    resume = parse_resume(resume_bytes)
    b = next(b for b in resume.blocks if "approval rates" in b.text)
    final = {b.id: b.text for b in resume.blocks}
    final[b.id] = b.text.replace("cross-functional teams", "multidisciplinary teams")
    result = render_comparison(resume_bytes, final)
    doc = Document(BytesIO(result))
    para = dict(document_paragraphs(doc))[b.location]
    assert para._p.xpath("./w:del/w:r/w:delText")
    assert "".join(p.text for p in para._p.xpath("./w:ins/w:r/w:t")) == "multidisciplinary"
    assert para.runs[0].text == "Led " and para.runs[0].bold
    assert "approval rates by 50%" in para.text  # Unchanged context stays outside revisions.
    assert len(doc.tables) == len(Document(BytesIO(resume_bytes)).tables)
    assert doc.sections[0].page_width == Document(BytesIO(resume_bytes)).sections[0].page_width
    with ZipFile(BytesIO(result)) as archive:
        xml = archive.read("word/settings.xml").decode()
        assert 'w:trackRevisions w:val="true"' in xml
        assert 'w:insDel="true"' in xml


def test_revision_ids_unique_across_tables_and_headers(resume_bytes):
    source = Document(BytesIO(resume_bytes))
    source.sections[0].header.paragraphs[0].text = "Old header"
    out = BytesIO()
    source.save(out)
    data = out.getvalue()
    resume = parse_resume(data)
    final = {b.id: b.text for b in resume.blocks}
    for b in resume.blocks:
        if b.text in ("Old header", "EXAMPLE COMPANY"):
            final[b.id] = b.text.replace("Old", "New").replace("EXAMPLE", "UPDATED")
    result = render_comparison(data, final)
    doc = Document(BytesIO(result))
    ids = [
        node.get(qn("w:id"))
        for _, p in document_paragraphs(doc)
        for node in p._p.xpath(".//w:ins|.//w:del")
    ]
    assert len(ids) == 4 and len(ids) == len(set(ids))
    for _, p in document_paragraphs(doc):
        for node in p._p.xpath(".//w:ins|.//w:del"):
            assert node.get(qn("w:author")) == "Resume Studio"
            assert node.get(qn("w:date")).endswith("Z")


def test_risk_comments_survive_accepting_or_rejecting(resume_bytes):
    resume = parse_resume(resume_bytes)
    block = next(b for b in resume.blocks if "approval rates" in b.text)
    final = {b.id: b.text for b in resume.blocks}
    final[block.id] = block.text.replace("supporting", "owning")
    risk = Risk(
        id="R-001",
        block_id=block.id,
        severity="suspected",
        origin="rewrite",
        quote="owning",
        source_text=block.text,
        reason="归因被加强",
        suggestion="核实",
    )
    result = render_comparison(resume_bytes, final, [risk])
    for accept in (True, False):
        doc = accepted_or_rejected(result, accept)
        assert "R-001｜待核实" in list(doc.comments)[0].text
        p = dict(document_paragraphs(doc))[block.location]
        assert p._p.xpath("./w:commentRangeStart")
        assert p._p.xpath("./w:commentRangeEnd")


def test_unchanged_paragraph_xml_stays_unchanged(resume_bytes):
    doc = Document(BytesIO(resume_bytes))
    expected = [p._p.xml for p in doc.paragraphs]
    final = {b.id: b.text for b in parse_resume(resume_bytes).blocks}
    actual = Document(BytesIO(render_comparison(resume_bytes, final)))
    assert not actual.element.xpath(".//w:ins|.//w:del")
    assert [p._p.xml for p in actual.paragraphs] == expected


def test_comparison_requires_exact_source_mapping(resume_bytes):
    with pytest.raises(RenderError):
        render_comparison(resume_bytes, {"wrong": "text"})


def test_saved_run_export_without_network_or_overwriting_input(tmp_path, resume_bytes):
    source = tmp_path / "source_resume.docx"
    source.write_bytes(resume_bytes)
    final = {b.id: b.text for b in parse_resume(resume_bytes).blocks}
    result = tmp_path / "result.json"
    result.write_text(json.dumps({"final_text": final, "risks": []}))
    original_result = result.read_bytes()
    path = export_saved_comparison(tmp_path)
    assert path.name == "resume_comparison.docx" and path.exists()
    assert source.read_bytes() == resume_bytes
    assert result.read_bytes() == original_result


def test_saved_real_runs_have_lossless_revision_views():
    root = Path(__file__).resolve().parents[1] / "outputs"
    files = list(root.glob("*/result.json"))
    if not files:
        pytest.skip("No private saved runs")
    for path in files:
        raw = json.loads(path.read_text())
        source = (path.parent / "source_resume.docx").read_bytes()
        render_comparison(source, raw["final_text"], [Risk.model_validate(r) for r in raw["risks"]])
