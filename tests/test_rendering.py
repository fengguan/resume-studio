from io import BytesIO
from zipfile import ZipFile

from docx import Document

from resume_studio.parsing import parse_resume
from resume_studio.rendering import render_resume
from resume_studio.schemas import Risk


def test_header_risk_highlights_source_with_valid_body_comment(resume_bytes):
    doc = Document(BytesIO(resume_bytes))
    doc.sections[0].header.paragraphs[0].text = "Certificate expected Aug 2020"
    out = BytesIO()
    doc.save(out)
    data = out.getvalue()
    resume = parse_resume(data)
    block = next(b for b in resume.blocks if b.location.startswith("section"))
    risk = Risk(
        id="R-001",
        block_id=block.id,
        severity="suspected",
        origin="source",
        quote="expected Aug 2020",
        source_text=block.text,
        reason="日期已过",
        suggestion="核实实际状态",
    )
    final = {b.id: b.text for b in resume.blocks}
    review = render_resume(data, final, [risk], annotate=True)
    with ZipFile(BytesIO(review)) as archive:
        assert "commentRangeStart" in archive.read("word/document.xml").decode()
        header = archive.read("word/header1.xml").decode()
        assert "highlight" in header and "commentRangeStart" not in header
    clean = render_resume(data, final, [], annotate=False)
    with ZipFile(BytesIO(clean)) as archive:
        assert "审阅提示" not in archive.read("word/document.xml").decode()


def test_multiple_overlapping_annotations_preserve_text_and_style(resume_bytes):
    resume = parse_resume(resume_bytes)
    block = next(b for b in resume.blocks if "approval rates" in b.text)
    risks = [
        Risk(
            id=f"R-00{i + 1}",
            block_id=block.id,
            severity="suspected",
            origin="source",
            quote=quote,
            source_text=block.text,
            reason="测试疑点",
            suggestion="核实",
        )
        for i, quote in enumerate(("approval rates by 50%", "50% and supporting $10B"))
    ]
    final = {b.id: b.text for b in resume.blocks}
    review = Document(BytesIO(render_resume(resume_bytes, final, risks, annotate=True)))
    para = next(p for p in review.paragraphs if "approval rates" in p.text)
    assert para.text == block.text
    assert para.runs[0].bold
    assert len(list(review.comments)) == 2
