import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from resume_studio.llm import ModelError
from resume_studio.schemas import Analysis, Audit, Draft, Mapping, Verdict


@pytest.fixture
def resume_bytes():
    doc = Document()
    doc.add_paragraph("TEST CANDIDATE")
    doc.add_paragraph("New York | person@example.com | (201) 555-0100")
    doc.add_paragraph("STRATEGY | ANALYTICS | INNOVATION")
    doc.add_paragraph(
        "Strategy leader with experience translating market research into executive recommendations and leading cross-functional projects across scientific and business teams."
    )
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "EXAMPLE COMPANY"
    table.cell(0, 1).text = "2020 - 2026"
    doc.add_paragraph("Project Lead (2020 - 2026)")
    p = doc.add_paragraph()
    p.add_run("Led ").bold = True
    p.add_run(
        "cross-functional teams to redesign the drug-evaluation system, increasing executive approval rates by 50% and supporting $10B in revenue potential."
    )
    doc.add_paragraph(
        "Evaluated 200+ early-stage innovation ideas and identified 30+ high-potential concepts through structured screening, strategic assessment, and prioritization."
    )
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("MBA, Example University")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def job_bytes():
    return json.dumps(
        {
            "job": {"title": "Strategy Manager", "company": "Example Employer"},
            "responsibilities": ["Lead portfolio prioritization", "Build financial models"],
            "qualifications": {"required": ["MBA or equivalent experience"]},
            "skills": ["DCF", "SQL"],
        }
    ).encode()


class FakeModel:
    """Deterministic test double, never presented as live AI output."""

    def __init__(self, changes=None, audit_fn=None, fail_audit=False):
        self.config = SimpleNamespace(provider="openai", model="TEST-DOUBLE")
        self.usage = []
        self.changes = changes or []
        self.audit_fn = audit_fn
        self.fail_audit = fail_audit
        self.payloads = []

    def generate(self, stage, system, payload, output_type):
        self.payloads.append((stage, payload))
        self.usage.append({"stage": stage, "tokens": {}})
        if stage == "analysis":
            return Analysis(
                positioning="优先突出已有组合评估经验。",
                priorities=["组合评估"],
                questions=["是否实际建立过 DCF 模型？"],
                mappings=[
                    Mapping(
                        requirement_id=r["id"],
                        status="missing",
                        evidence_ids=[],
                        reason="测试用对照结论",
                    )
                    for r in payload["job"]["requirements"]
                ],
            )
        if stage == "tailoring":
            return Draft(changes=self.changes)
        if self.fail_audit:
            raise ModelError("测试：审校超时")
        if self.audit_fn:
            return self.audit_fn(payload)
        return Audit(
            verdicts=[
                Verdict(
                    block_id=b,
                    severity="clear",
                    origin="source",
                    quote="",
                    reason="",
                    suggestion="",
                )
                for b in payload["review_block_ids"]
            ]
        )


@pytest.fixture
def fake_model():
    return FakeModel
